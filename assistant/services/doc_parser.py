import os
import PyPDF2
from docx import Document as DocxDocument
from pptx import Presentation

def extract_text_from_file(file_path, file_type):
    """
    Extract text content from uploaded file based on extension (.pdf, .docx, .pptx, .txt).
    Returns clean string text.
    """
    file_type = file_type.lower().replace('.', '')
    text = ""

    try:
        if file_type == 'pdf':
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                pages_text = []
                for i, page in enumerate(reader.pages):
                    page_content = page.extract_text() or ""
                    if page_content.strip():
                        pages_text.append(f"[Page {i+1}]\n{page_content.strip()}")
                text = "\n\n".join(pages_text)

        elif file_type in ['docx', 'doc']:
            doc = DocxDocument(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)

        elif file_type in ['pptx', 'ppt']:
            prs = Presentation(file_path)
            slides_text = []
            for i, slide in enumerate(prs.slides):
                slide_content = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_content.append(shape.text.strip())
                if slide_content:
                    slides_text.append(f"[Slide {i+1}]\n" + "\n".join(slide_content))
            text = "\n\n".join(slides_text)

        elif file_type in ['txt', 'md']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()

        elif file_type in ['png', 'jpg', 'jpeg', 'webp', 'bmp']:
            # Multimodal Vision Extraction via PIL & Gemini Vision
            text = extract_multimodal_image_insights(file_path)

        else:
            text = f"Unsupported file type: {file_type}"

    except Exception as e:
        print(f"Error extracting text from {file_path}: {e}")
        text = f"Error processing file: {str(e)}"

    return clean_text(text)

def clean_text(text):
    if not text:
        return ""
    # Basic text normalization
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)

def extract_multimodal_image_insights(image_path):
    """
    Extracts deep visual semantic context, OCR text, data charts, and architecture diagrams using Gemini Vision.
    """
    try:
        from PIL import Image
        import google.generativeai as genai
        import os

        img = Image.open(image_path)
        api_key = os.getenv('GEMINI_API_KEY', '').strip()

        if api_key:
            genai.configure(api_key=api_key)
            prompt = """
You are ScholarPulse AI's Multimodal Visual Perception Engine.
Analyze this academic image / diagram / flowchart / architecture chart / research figure in detail.

Provide a comprehensive extraction in the following format:
[VISUAL ENTITY & DIAGRAM TYPE]
- Identify if this is a System Architecture, Flowchart, Scatter Plot, Bar Chart, Mathematical Formula, or UI Mockup.

[TRANSCRIPTION & LABELS]
- Transcribe all text, numbers, formulas, and labels present in the image.

[SYSTEM ARCHITECTURE & DATA FLOW BREAKDOWN]
- Detail each component, input streams, processing layers, and output outcomes.

[RESEARCH IMPLICATIONS & CORE TAKEAWAYS]
- Explain what this figure contributes to the research project or technical study.
"""
            model_names = ['gemini-2.0-flash', 'gemini-1.5-flash-latest', 'gemini-1.5-flash', 'gemini-pro-vision']
            for m_name in model_names:
                try:
                    model = genai.GenerativeModel(m_name)
                    res = model.generate_content([prompt, img])
                    if res and res.text:
                        return f"[Multimodal Visual Analysis: {os.path.basename(image_path)}]\n\n" + res.text
                except Exception:
                    continue

        # Fallback if no API key or vision call fails
        width, height = img.size
        format_name = img.format or 'IMAGE'
        return f"[Image Document: {os.path.basename(image_path)}]\nResolution: {width}x{height}px\nFormat: {format_name}\nContains visual research diagram / architectural layout."

    except Exception as e:
        print(f"Error in extract_multimodal_image_insights: {e}")
        return f"[Image Document: {os.path.basename(image_path)}]\nError analyzing visual contents: {str(e)}"

def fetch_and_parse_url(url, custom_title=None):
    """
    Downloads content from a PDF or web page URL and extracts text.
    Returns dict with title, filename, extracted_text, file_bytes, extension.
    """
    import tempfile
    import urllib.parse
    import requests

    url = url.strip()
    if not url.startswith(('http://', 'https://')):
        url = 'https://' + url

    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) ScholarPulse-AI/1.0 (Academic Research Assistant)'
    }

    try:
        response = requests.get(url, headers=headers, timeout=20, stream=True)
        response.raise_for_status()

        content_type = response.headers.get('Content-Type', '').lower()
        parsed_url = urllib.parse.urlparse(url)
        url_filename = os.path.basename(parsed_url.path) or 'web_document'

        # Check if URL or Content-Type is PDF
        if 'application/pdf' in content_type or url.lower().endswith('.pdf'):
            ext = '.pdf'
            filename = url_filename if url_filename.endswith('.pdf') else f"{url_filename}.pdf"
            pdf_bytes = response.content

            # Save temporarily to parse PDF
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
                tmp.write(pdf_bytes)
                tmp_path = tmp.name

            try:
                extracted_text = extract_text_from_file(tmp_path, 'pdf')
            finally:
                if os.path.exists(tmp_path):
                    try:
                        os.remove(tmp_path)
                    except Exception:
                        pass

            derived_title = custom_title or os.path.splitext(filename)[0].replace('_', ' ').replace('-', ' ').title()
            return {
                'title': derived_title,
                'filename': filename,
                'extracted_text': extracted_text,
                'file_bytes': pdf_bytes,
                'extension': '.pdf'
            }

        else:
            # HTML Webpage Parsing
            html_content = response.text
            extracted_title = ""
            main_text = ""

            try:
                from bs4 import BeautifulSoup
                soup = BeautifulSoup(html_content, 'html.parser')

                # Extract title
                title_tag = soup.find('title')
                if title_tag and title_tag.string:
                    extracted_title = title_tag.string.strip()

                # Remove script, style, nav, footer, header elements
                for element in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
                    element.decompose()

                lines = []
                for elem in soup.find_all(['h1', 'h2', 'h3', 'h4', 'p', 'li']):
                    txt = elem.get_text().strip()
                    if txt:
                        lines.append(txt)

                main_text = "\n\n".join(lines)
            except Exception:
                import re
                clean_html = re.sub(r'<script.*?>.*?</script>', '', html_content, flags=re.DOTALL)
                clean_html = re.sub(r'<style.*?>.*?</style>', '', clean_html, flags=re.DOTALL)
                clean_html = re.sub(r'<.*?>', ' ', clean_html)
                main_text = "\n".join([line.strip() for line in clean_html.splitlines() if line.strip()])

            derived_title = custom_title or extracted_title or parsed_url.netloc or "Web Article Document"
            filename = f"web_{parsed_url.netloc.replace('.', '_')}.txt"
            file_bytes = main_text.encode('utf-8')

            return {
                'title': derived_title,
                'filename': filename,
                'extracted_text': clean_text(main_text),
                'file_bytes': file_bytes,
                'extension': '.txt'
            }

    except Exception as e:
        raise ValueError(f"Failed to fetch content from URL '{url}': {str(e)}")


